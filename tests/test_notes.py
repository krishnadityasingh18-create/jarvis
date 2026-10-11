import json
import os

import numpy as np

from jarvis import db, notes

WORDS = ["buck", "boost", "inductor", "physics", "exam"]
SYNONYMS = {"step-down": "buck", "step-up": "boost"}


def fake_embed(texts, kind):
    """One dimension per keyword; synonyms share a dimension (so vectors can find what words can't)."""
    out = []
    for t in texts:
        t = t.lower()
        for a, b in SYNONYMS.items():
            t = t.replace(a, b)
        v = np.array([1.0 if w in t else 0.0 for w in WORDS] + [0.01], np.float32)
        out.append(v / np.linalg.norm(v))
    return np.array(out)


def no_embed(texts, kind):
    return None


def make_pdf(path, pages):
    import pymupdf
    d = pymupdf.open()
    for text in pages:
        d.new_page().insert_text((72, 72), text)
    d.save(path)


def make_docs(root):
    (root / "notes").mkdir(parents=True)
    make_pdf(root / "notes" / "power.pdf", ["Buck converter basics", "Boost converter page"])
    import docx
    d = docx.Document()
    d.add_paragraph("Physics exam on Monday")
    d.save(root / "notes" / "plan.docx")
    import pptx
    p = pptx.Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Inductor slide"
    p.save(root / "notes" / "deck.pptx")
    (root / "notes" / "nb.ipynb").write_text(json.dumps(
        {"cells": [{"cell_type": "markdown", "source": ["# Physics lab\n"]},
                   {"cell_type": "code", "source": ["print('boost')"]}]}))
    (root / "notes" / "todo.md").write_text("exam revision list")


def test_extract_pdf_docx_pptx_ipynb_txt(tmp_path):
    make_docs(tmp_path)
    n = tmp_path / "notes"
    assert [(p, t.strip()) for p, t in notes.extract(n / "power.pdf")] == \
        [(1, "Buck converter basics"), (2, "Boost converter page")]
    assert "Physics exam on Monday" in notes.extract(n / "plan.docx")[0][1]
    assert notes.extract(n / "deck.pptx")[0] == (1, "Inductor slide")
    nb = notes.extract(n / "nb.ipynb")[0][1]
    assert "# Physics lab" in nb and "print('boost')" in nb
    assert notes.extract(n / "todo.md") == [(0, "exam revision list")]


def test_chunk_overlap_and_pages():
    text = "".join(f"{i:04d} " for i in range(500))      # 2500 chars
    out = notes.chunk([(3, text), (4, "short")], size=1000, overlap=150)
    assert [p for p, _ in out] == [3, 3, 3, 4]
    assert out[0][1][-150:] == out[1][1][:150] and out[-1] == (4, "short")


def index(tmp_path, cfg, embed=fake_embed, exclude=()):
    return notes.index_contents(cfg.db_path, [tmp_path / "home"], list(exclude), embed=embed)


def test_index_incremental(tmp_path, cfg, conn):
    home = tmp_path / "home"
    make_docs(home)
    s = index(tmp_path, cfg)
    assert s["added"] == 5 and s["failed"] == 0 and s["embedded"] > 0
    assert index(tmp_path, cfg)["added"] == 0                      # unchanged: skipped
    todo = home / "notes" / "todo.md"
    todo.write_text("new boost text")
    os.utime(todo, (2_000_000_000, 2_000_000_000))
    (home / "notes" / "plan.docx").unlink()
    s = index(tmp_path, cfg)
    assert s["updated"] == 1 and s["removed"] == 1
    texts = [r[0] for r in conn.execute("SELECT text FROM chunks WHERE path=?", (str(todo),))]
    assert texts == ["new boost text"]
    assert not conn.execute("SELECT 1 FROM chunks WHERE path LIKE '%plan.docx'").fetchone()


def test_index_survives_bad_pdf(tmp_path, cfg, conn):
    make_docs(tmp_path / "home")
    (tmp_path / "home" / "notes" / "broken.pdf").write_bytes(b"not a pdf at all")
    s = index(tmp_path, cfg)
    assert s["failed"] == 1 and s["added"] == 5
    err = conn.execute("SELECT error FROM docs WHERE path LIKE '%broken.pdf'").fetchone()[0]
    assert err


def test_index_skips_content_exclude(tmp_path, cfg, conn):
    make_docs(tmp_path / "home")
    (tmp_path / "home" / "My project").mkdir()
    (tmp_path / "home" / "My project" / "readme.md").write_text("unity boost")
    index(tmp_path, cfg, exclude=["My project"])
    assert not conn.execute("SELECT 1 FROM docs WHERE path LIKE '%readme.md'").fetchone()


def test_search_fuses_fts_and_vectors(tmp_path, cfg, conn):
    make_docs(tmp_path / "home")
    index(tmp_path, cfg)
    hits = notes.search(conn, "step-down converter", k=3, embed=fake_embed)
    assert hits[0]["path"].endswith("power.pdf") and hits[0]["page"] == 1   # vector + words
    hits = notes.search(conn, "revision", k=3, embed=fake_embed)                # words only
    assert hits[0]["path"].endswith("todo.md")


def test_search_without_vectors(tmp_path, cfg, conn):
    make_docs(tmp_path / "home")
    index(tmp_path, cfg, embed=no_embed)
    assert conn.execute("SELECT count(*) FROM chunks WHERE vec IS NULL").fetchone()[0] == 5 + 1
    hits = notes.search(conn, "inductor", k=3, embed=no_embed)
    assert hits[0]["path"].endswith("deck.pptx") and hits[0]["page"] == 1
