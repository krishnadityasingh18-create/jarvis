"""The content index: text extraction, chunks, full-text + vector search. Local only."""
import json
import os
import re
from pathlib import Path

import numpy as np

from jarvis import db

# Notes and code. No .json/.csv/.css: in practice they're tool caches and datasets, not notes.
# No .html either: on this laptop it's generated library docs (Doxygen), not notes.
CODE_EXT = {".py", ".c", ".cpp", ".h", ".hpp", ".ino", ".m", ".java", ".js", ".ts", ".sql"}
TEXT_EXT = {".md", ".txt", ".rst", ".tex", ".yaml", ".yml", ".toml", ".ini"} | CODE_EXT
MAX_CODE = 200 * 2**10   # bigger code files are bundles or generated, not something you wrote
DOC_EXT = TEXT_EXT | {".pdf", ".docx", ".pptx", ".ipynb"}
MAX_PDF = 50 * 2**20
MAX_OTHER = 2 * 2**20
MAX_PAGES = 300
EMBED_MODEL = "nomic-embed-text"
OLLAMA = "http://localhost:11434/v1"
BATCH = 64


def extract(path: Path) -> list[tuple[int, str]]:
    """(page, text) pairs; page is 1-based for PDFs and slides, 0 otherwise."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(path) as d:
            if d.needs_pass:
                raise ValueError("the PDF is password-protected")
            return [(i + 1, p.get_text()) for i, p in enumerate(d) if i < MAX_PAGES]
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
        return [(0, "\n".join([p.text for p in d.paragraphs] + cells))]
    if ext == ".pptx":
        import pptx
        return [(i + 1, "\n".join(t for s in slide.shapes
                                  if s.has_text_frame and (t := s.text_frame.text.strip())))
                for i, slide in enumerate(pptx.Presentation(path).slides) if i < MAX_PAGES]
    if ext == ".ipynb":
        cells = json.loads(path.read_text(encoding="utf-8", errors="replace"))["cells"]
        return [(0, "\n\n".join("".join(c.get("source", [])) for c in cells))]
    return [(0, path.read_text(encoding="utf-8", errors="replace"))]


def chunk(pages, size: int = 1000, overlap: int = 150) -> list[tuple[int, str]]:
    out = []
    for page, text in pages:
        text = re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()
        i = 0
        while i < len(text):
            if text[i:i + size].strip():
                out.append((page, text[i:i + size]))
            if i + size >= len(text):
                break
            i += size - overlap
    return out


def embed(texts: list[str], kind: str) -> np.ndarray | None:
    """Unit-length vectors from Ollama, or None when it isn't reachable."""
    import openai
    prefix = "search_query: " if kind == "query" else "search_document: "
    try:
        client = openai.OpenAI(base_url=OLLAMA, api_key="ollama", max_retries=0, timeout=120)
        r = client.embeddings.create(model=EMBED_MODEL, input=[prefix + t for t in texts])
    except Exception:
        return None
    v = np.array([d.embedding for d in r.data], np.float32)
    return v / np.maximum(np.linalg.norm(v, axis=1, keepdims=True), 1e-9)


def _minified(path: str) -> bool:
    """Generated or minified text (one enormous line) is noise for answers."""
    try:
        with open(path, encoding="utf-8", errors="replace") as f:
            sample = f.read(65536)
    except OSError:
        return True
    lines = sample.splitlines() or [""]
    return max(map(len, lines)) > 3000 or len(sample) / len(lines) > 250


def _walk(roots, exclude, skip):
    names = set(exclude)
    skip = [Path(s).resolve() for s in skip]
    for root in roots:
        for dirpath, dirnames, filenames in os.walk(root):
            if any(Path(dirpath).resolve().is_relative_to(s) for s in skip):
                dirnames[:] = []
                continue
            # hidden folders hold tool data (caches, snapshots), not notes
            dirnames[:] = [d for d in dirnames if d not in names and not d.startswith(".")]
            for f in filenames:
                ext = os.path.splitext(f)[1].lower()
                if ext not in DOC_EXT or f.startswith(("~$", ".~")):   # Office lock files
                    continue
                p = os.path.join(dirpath, f)
                try:
                    st = os.stat(p)
                except OSError:
                    continue
                cap = MAX_PDF if ext == ".pdf" else MAX_CODE if ext in CODE_EXT else MAX_OTHER
                if st.st_size > cap:
                    continue
                if ext in TEXT_EXT and _minified(p):
                    continue
                yield p, st.st_mtime, st.st_size


def index_contents(db_path, roots, exclude, embed=embed, skip=()) -> dict:
    """Bring the content index up to date: re-read changed files, drop removed ones, then
    fill in missing vectors. Each file commits on its own so the lock is held briefly."""
    stats = dict(added=0, updated=0, removed=0, failed=0, embedded=0)
    c = db.connect(db_path)
    try:
        known = {r[0]: (r[1], r[2]) for r in c.execute("SELECT path, mtime, size FROM docs")}
        seen = set()
        for path, mtime, size in _walk(roots, exclude, skip):
            seen.add(path)
            if known.get(path) == (mtime, size):
                continue
            try:
                pieces, error = chunk(extract(Path(path))), None
            except Exception as e:
                pieces, error = [], f"{type(e).__name__}: {e}"[:200]
            c.execute("BEGIN")
            c.execute("DELETE FROM chunks WHERE path=?", (path,))
            c.executemany("INSERT INTO chunks(path, page, text) VALUES(?,?,?)",
                          [(path, p, t) for p, t in pieces])
            c.execute("INSERT OR REPLACE INTO docs(path, mtime, size, chunks, error) "
                      "VALUES(?,?,?,?,?)", (path, mtime, size, len(pieces), error))
            c.execute("COMMIT")
            stats["failed" if error else "updated" if path in known else "added"] += 1
        for path in set(known) - seen:
            c.execute("DELETE FROM chunks WHERE path=?", (path,))
            c.execute("DELETE FROM docs WHERE path=?", (path,))
            stats["removed"] += 1
        while rows := c.execute("SELECT id, text FROM chunks WHERE vec IS NULL LIMIT ?",
                                (BATCH,)).fetchall():
            v = embed([r[1] for r in rows], "document")
            if v is None:
                break  # Ollama down: full-text search still works; a later pass fills these
            c.execute("BEGIN")
            c.executemany("UPDATE chunks SET vec=? WHERE id=?",
                          [(v[i].tobytes(), r[0]) for i, r in enumerate(rows)])
            c.execute("COMMIT")
            stats["embedded"] += len(rows)
    finally:
        c.close()
    return stats


_matrix: dict = {}


def _vectors(conn):
    """All chunk vectors as one matrix, reloaded only when the set of vectors changes."""
    key = (id(conn),) + tuple(conn.execute(
        "SELECT count(*), coalesce(max(id), 0) FROM chunks WHERE vec IS NOT NULL").fetchone())
    if _matrix.get("key") != key:
        rows = conn.execute("SELECT id, vec FROM chunks WHERE vec IS NOT NULL").fetchall()
        _matrix.update(key=key, ids=np.array([r[0] for r in rows]),
                       m=np.array([np.frombuffer(r[1], np.float32) for r in rows])
                       if rows else None)
    return _matrix["ids"], _matrix["m"]


def search(conn, query: str, k: int = 6, embed=embed) -> list[dict]:
    """Reciprocal-rank fusion of full-text (BM25) and vector results."""
    ranked = []
    words = re.findall(r"\w+", query.lower())
    if words:
        ranked.append([r[0] for r in conn.execute(
            "SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ? ORDER BY rank LIMIT 30",
            (" OR ".join(f'"{w}"' for w in words),))])
    ids, m = _vectors(conn)
    if m is not None and (q := embed([query], "query")) is not None:
        ranked.append([int(ids[i]) for i in np.argsort(-(m @ q[0]))[:30]])
    score: dict[int, float] = {}
    for lst in ranked:
        for rank, cid in enumerate(lst):
            score[cid] = score.get(cid, 0) + 1 / (60 + rank)
    out = []
    for cid in sorted(score, key=score.get, reverse=True)[:k]:
        path, page, text = conn.execute("SELECT path, page, text FROM chunks WHERE id=?",
                                        (cid,)).fetchone()
        out.append({"path": path, "page": page, "text": text, "score": round(score[cid], 4)})
    return out
